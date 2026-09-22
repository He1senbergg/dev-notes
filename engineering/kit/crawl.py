
import csv
import os
import re
import time
import base64
import json
import sys
import traceback
from urllib.parse import urlsplit
from datetime import datetime, timedelta

from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from DrissionPage import ChromiumPage, ChromiumOptions
from DrissionPage.errors import ElementNotFoundError

TARGET_URL = "https://hr.huawei.com/orgarchive/index.html#/?orgcode=075774"
# CSV_FILE = r"D:\Desktop\24hCrawl\employee_stats.csv"
EXCEL_FILE = "employee_stats.xlsx"
INTERVAL_SECONDS = 86400
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BROWSER_USER_DATA = os.path.join(BASE_DIR, 'browser_profile')
BROWSER_PORT = 9333
LOGIN_TIMEOUT = 180
W3_URL = 'https://w3.huawei.com/next/indexa.html'
W3_LOGIN_NAME = 'Chen Fumin'

SNAPSHOT_DIR = "snapshots"
CARD_TITLE = "ICT BG"
LOG_DIR = r"D:\Desktop\24hCrawl\logs"
LOG_FILE_PATH = None

def init_log_file() -> str:
    """
    初始化日志文件路径（自动编号）
    """
    global LOG_FILE_PATH

    if not os.path.exists(LOG_DIR):
        os.makedirs(LOG_DIR)

    files = os.listdir(LOG_DIR)

    max_index = 0

    for file in files:
        match = re.match(r"log_(\d+)\.txt", file)
        if match:
            index = int(match.group(1))
            max_index = max(max_index, index)

    new_index = max_index + 1

    LOG_FILE_PATH = os.path.join(LOG_DIR, f"log_{new_index}.txt")

    return LOG_FILE_PATH


def log_info(message: str) -> None:
    """
    写日志到文件
    """
    global LOG_FILE_PATH

    if LOG_FILE_PATH is None:
        init_log_file()

    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    log_line = f"[log_info] {timestamp} {message}\n"

    print(f"[log_info] {timestamp} {message}")
    with open(LOG_FILE_PATH, "a", encoding="utf-8") as f:
        f.write(log_line)


# def create_page() -> ChromiumPage:
#     """
#     创建浏览器页面对象
#     """
#     log_info("[create_page] 开始创建浏览器页面对象")

#     co = ChromiumOptions()
#     co.auto_port()

#     page = ChromiumPage(co)

#     try:
#         page.set.window.max()
#         log_info("[create_page] 浏览器窗口已最大化")
#     except Exception as e:
#         log_info(f"[create_page] 浏览器窗口最大化失败: {e}")

#     log_info("[create_page] 浏览器页面对象创建成功")
#     return page


def create_page() -> ChromiumPage:
    """
    创建浏览器页面对象
    """
    log_info("[create_page] 开始创建浏览器页面对象")

    co = ChromiumOptions(read_file=False)
    # 固定端口和专用配置目录，跨次运行保留登录状态。
    co.set_local_port(BROWSER_PORT)
    co.set_user_data_path(BROWSER_USER_DATA)

    page = ChromiumPage(co)

    log_info("[create_page] 浏览器页面对象创建成功")
    return page


def parse_number(text: str) -> int:
    """
    从文本中提取数字
    """
    log_info(f"开始解析数字，原始文本: {text}")

    match = re.search(r"\d+", text)
    if not match:
        raise ValueError(f"未从文本中提取到数字，原始文本: {text}")

    value = int(match.group())
    log_info(f"数字解析成功: {value}")
    return value


def get_count_by_label(page: ChromiumPage, label_text: str) -> int:
    """
    根据标签名获取对应数量
    """
    log_info(f"开始抓取字段: {label_text}")

    xpath = f"xpath://div[contains(@class, 'content')][.//p[normalize-space()='{label_text}']]/p[2]"
    ele = page.ele(xpath, timeout=20)

    if ele is None:
        raise ValueError(f"未找到字段元素: {label_text}")

    raw_text = ele.text.strip()
    log_info(f"字段 {label_text} 抓取到原始文本: {raw_text}")

    value = parse_number(raw_text)
    log_info(f"字段 {label_text} 抓取完成，结果: {value}")
    return value


def save_to_csv(file_path: str, row: list) -> None:
    """
    保存数据到 CSV
    """
    file_exists = os.path.exists(file_path)

    with open(file_path, mode="a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)

        if not file_exists:
            writer.writerow(["日期", "星期", "华为员工", "外包员工"])

        writer.writerow(row)

    log_info(f"数据写入 CSV 成功: {row}")


def sanitize_filename(name: str) -> str:
    """
    将文件名中的非法字符替换掉
    """
    invalid_chars = ['\\', '/', ':', '*', '?', '"', '<', '>', '|']
    result = name
    for ch in invalid_chars:
        result = result.replace(ch, "_")
    return result


def wait_card_ready(page: ChromiumPage, card_ele, timeout: int = 10) -> None:
    """
    等待原始卡片渲染完成，避免头像未加载、卡片高度未稳定
    """
    log_info("[wait_card_ready] 开始等待原始卡片渲染稳定")

    start_time = time.time()
    last_height = None
    stable_count = 0

    while True:
        try:
            js = """
            const ele = arguments[0];
            const rect = ele.getBoundingClientRect();
            const imgs = ele.querySelectorAll('img');

            let all_imgs_ready = true;
            for (const img of imgs) {
                if (!img.complete || img.naturalWidth <= 0) {
                    all_imgs_ready = false;
                    break;
                }
            }

            return JSON.stringify({
                height: rect.height,
                width: rect.width,
                all_imgs_ready: all_imgs_ready
            });
            """

            result = page.run_js(js, card_ele)
            data = json.loads(result)

            height = round(float(data["height"]), 2)
            width = round(float(data["width"]), 2)
            all_imgs_ready = data["all_imgs_ready"]

            log_info(
                f"[wait_card_ready] 当前状态: width={width}, height={height}, "
                f"all_imgs_ready={all_imgs_ready}, stable_count={stable_count}"
            )

            time.sleep(2)  # 等图片加载的间隔
            if all_imgs_ready:
                if last_height is not None and abs(height - last_height) < 0.5:
                    stable_count += 1
                else:
                    stable_count = 0

                last_height = height

                if stable_count >= 2:
                    log_info("[wait_card_ready] 原始卡片已渲染稳定")
                    return

        except Exception as e:
            log_info(f"[wait_card_ready] 检测过程中异常: {e}")

        if time.time() - start_time > timeout:
            log_info("[wait_card_ready] 等待超时，继续执行当前截图流程")
            return

        time.sleep(0.5)


def get_card_bbox_by_js(page: ChromiumPage, card_ele) -> dict:
    """
    使用 JS 获取元素在整个页面中的真实坐标和尺寸
    """
    log_info("[get_card_bbox_by_js] 开始获取卡片坐标")

    js = """
    const ele = arguments[0];
    const rect = ele.getBoundingClientRect();
    return JSON.stringify({
        x: rect.left + window.scrollX,
        y: rect.top + window.scrollY,
        width: rect.width,
        height: rect.height,
        dpr: window.devicePixelRatio
    });
    """

    result = page.run_js(js, card_ele)
    data = json.loads(result)

    log_info(
        f"[get_card_bbox_by_js] 获取成功: "
        f"x={data['x']}, y={data['y']}, width={data['width']}, "
        f"height={data['height']}, dpr={data['dpr']}"
    )
    return data


def capture_card_by_cdp(page: ChromiumPage, bbox: dict, image_path: str) -> None:
    """
    使用 Chrome DevTools Protocol 按坐标截取页面区域
    """
    log_info("[capture_card_by_cdp] 开始使用 CDP 截图")

    margin = 0

    clip_x = max(0, float(bbox["x"]) - margin)
    clip_y = max(0, float(bbox["y"]) - margin)
    clip_width = float(bbox["width"]) + margin * 2
    clip_height = float(bbox["height"]) + margin * 2

    log_info(
        f"[capture_card_by_cdp] 截图区域: "
        f"x={clip_x}, y={clip_y}, width={clip_width}, height={clip_height}"
    )

    try:
        page.run_cdp("Page.enable")
        log_info("[capture_card_by_cdp] Page.enable 执行成功")
    except Exception as e:
        log_info(f"[capture_card_by_cdp] Page.enable 执行失败: {e}")

    result = page.run_cdp(
        "Page.captureScreenshot",
        format="png",
        fromSurface=True,
        captureBeyondViewport=True,
        clip={
            "x": clip_x,
            "y": clip_y,
            "width": clip_width,
            "height": clip_height,
            "scale": 1
        }
    )

    image_data = base64.b64decode(result["data"])

    with open(image_path, "wb") as f:
        f.write(image_data)

    log_info(f"[capture_card_by_cdp] CDP截图保存成功，路径={image_path}")


def find_visible_card(page: ChromiumPage, card_title: str):
    """
    查找真正可见且尺寸正常的目标卡片
    """
    log_info(f"[find_visible_card] 开始查找标题为 {card_title} 的可见卡片")

    xpath = (
        f"xpath://div[contains(@class,'infoCard')]"
        f"[.//div[contains(@class,'title') and normalize-space()='{card_title}']]"
    )

    cards = page.eles(xpath, timeout=20)

    if not cards:
        raise ValueError(f"[find_visible_card] 未找到任何标题为 {card_title} 的卡片")

    log_info(f"[find_visible_card] 共找到 {len(cards)} 个候选卡片")

    for index, card in enumerate(cards, start=1):
        try:
            is_displayed = card.states.is_displayed
        except Exception:
            is_displayed = False

        try:
            rect = card.rect
            width = rect.size[0]
            height = rect.size[1]
        except Exception:
            width = 0
            height = 0

        log_info(
            f"[find_visible_card] 检查第 {index} 个卡片: "
            f"is_displayed={is_displayed}, width={width}, height={height}"
        )

        if is_displayed and width > 50 and height > 50:
            log_info(f"[find_visible_card] 选中第 {index} 个可见卡片")
            return card

    raise RuntimeError(f"[find_visible_card] 找到了卡片，但没有可见且尺寸正常的目标卡片，标题={card_title}")


def clone_card_to_fixed_layer(page: ChromiumPage, card_ele) -> None:
    """
    将目标卡片克隆到页面左上角固定层
    """
    log_info("[clone_card_to_fixed_layer] 开始克隆卡片到固定层")

    js = """
    const ele = arguments[0];

    const oldWrap = document.getElementById('__chatgpt_capture_wrap__');
    if (oldWrap) {
        oldWrap.remove();
    }

    const wrap = document.createElement('div');
    wrap.id = '__chatgpt_capture_wrap__';
    wrap.style.position = 'fixed';
    wrap.style.left = '20px';
    wrap.style.top = '20px';
    wrap.style.zIndex = '2147483647';
    wrap.style.background = '#ffffff';
    wrap.style.padding = '12px';
    wrap.style.border = '1px solid #ddd';
    wrap.style.boxShadow = '0 2px 12px rgba(0,0,0,0.15)';
    wrap.style.overflow = 'visible';

    const clone = ele.cloneNode(true);
    clone.id = '__chatgpt_capture_clone__';

    clone.style.margin = '0';
    clone.style.transform = 'none';
    clone.style.position = 'relative';
    clone.style.left = '0';
    clone.style.top = '0';

    wrap.appendChild(clone);
    document.body.appendChild(wrap);
    """

    page.run_js(js, card_ele)
    log_info("[clone_card_to_fixed_layer] 克隆卡片完成")


def get_fixed_clone_bbox(page: ChromiumPage, timeout: int = 10) -> dict:
    """
    获取固定层克隆卡片的页面坐标 bbox
    注意：截图使用的是页面坐标，因此这里要加上 scrollX / scrollY
    """
    log_info("[get_fixed_clone_bbox] 开始获取固定层克隆卡片 bbox")

    start_time = time.time()

    while True:
        try:
            js = """
            const wrap = document.getElementById('__chatgpt_capture_wrap__');
            if (!wrap) {
                return JSON.stringify({
                    ok: false,
                    reason: 'wrap_not_found'
                });
            }

            const rect = wrap.getBoundingClientRect();

            return JSON.stringify({
                ok: true,
                x: rect.left + window.scrollX,
                y: rect.top + window.scrollY,
                width: rect.width,
                height: rect.height,
                viewport_x: rect.left,
                viewport_y: rect.top,
                scroll_x: window.scrollX,
                scroll_y: window.scrollY,
                dpr: window.devicePixelRatio
            });
            """

            result = page.run_js(js)
            data = json.loads(result)

            if data.get("ok"):
                log_info(
                    f"[get_fixed_clone_bbox] 获取成功: "
                    f"x={data['x']}, y={data['y']}, width={data['width']}, height={data['height']}, "
                    f"viewport_x={data['viewport_x']}, viewport_y={data['viewport_y']}, "
                    f"scroll_x={data['scroll_x']}, scroll_y={data['scroll_y']}, dpr={data['dpr']}"
                )
                return data

            log_info(f"[get_fixed_clone_bbox] 固定层未找到: {data}")

        except Exception as e:
            log_info(f"[get_fixed_clone_bbox] 获取 bbox 异常: {e}")

        if time.time() - start_time > timeout:
            raise TimeoutError("[get_fixed_clone_bbox] 获取固定层 bbox 超时")

        time.sleep(0.5)


def trim_white_border(image_path: str, color_tolerance: int = 18, edge_tolerance: float = 0.97) -> None:
    """
    按图片四角背景色自动裁掉四周接近背景色的边缘
    更适合卡片类截图，比单纯按纯白裁剪更紧
    """
    log_info(f"[trim_white_border] 开始裁剪白边: {image_path}")

    img = Image.open(image_path).convert("RGB")
    width, height = img.size
    pixels = img.load()

    # 取四个角的颜色，避免只取一个角不稳
    corner_colors = [
        pixels[0, 0],
        pixels[width - 1, 0],
        pixels[0, height - 1],
        pixels[width - 1, height - 1],
    ]

    bg_r = sum(c[0] for c in corner_colors) // 4
    bg_g = sum(c[1] for c in corner_colors) // 4
    bg_b = sum(c[2] for c in corner_colors) // 4

    log_info(
        f"[trim_white_border] 背景参考色: r={bg_r}, g={bg_g}, b={bg_b}, "
        f"color_tolerance={color_tolerance}, edge_tolerance={edge_tolerance}"
    )

    def is_near_bg(r: int, g: int, b: int) -> bool:
        return (
            abs(r - bg_r) <= color_tolerance and
            abs(g - bg_g) <= color_tolerance and
            abs(b - bg_b) <= color_tolerance
        )

    def col_is_bg_like(x: int) -> bool:
        match_count = 0
        for y in range(height):
            r, g, b = pixels[x, y]
            if is_near_bg(r, g, b):
                match_count += 1
        return match_count / height >= edge_tolerance

    def row_is_bg_like(y: int) -> bool:
        match_count = 0
        for x in range(width):
            r, g, b = pixels[x, y]
            if is_near_bg(r, g, b):
                match_count += 1
        return match_count / width >= edge_tolerance

    left = 0
    right = width - 1
    top = 0
    bottom = height - 1

    while left < width and col_is_bg_like(left):
        left += 1

    while right >= 0 and col_is_bg_like(right):
        right -= 1

    while top < height and row_is_bg_like(top):
        top += 1

    while bottom >= 0 and row_is_bg_like(bottom):
        bottom -= 1

    if left > right or top > bottom:
        log_info("[trim_white_border] 未找到有效内容区域，跳过裁剪")
        return

    # 再轻微内收 1 像素，让边更紧一点
    shrink = 1
    left = min(left + shrink, right)
    top = min(top + shrink, bottom)
    right = max(left, right - shrink)
    bottom = max(top, bottom - shrink)

    cropped = img.crop((left, top, right + 1, bottom + 1))
    cropped.save(image_path)

    log_info(
        f"[trim_white_border] 裁剪完成: "
        f"left={left}, top={top}, right={right}, bottom={bottom}"
    )


def remove_fixed_clone_layer(page: ChromiumPage) -> None:
    """
    删除临时固定克隆层
    """
    log_info("[remove_fixed_clone_layer] 开始清理固定克隆层")

    js = """
    const oldWrap = document.getElementById('__chatgpt_capture_wrap__');
    if (oldWrap) {
        oldWrap.remove();
    }
    """

    try:
        page.run_js(js)
        log_info("[remove_fixed_clone_layer] 固定克隆层清理完成")
    except Exception as e:
        log_info(f"[remove_fixed_clone_layer] 清理固定克隆层失败: {e}")


def save_card_screenshot(page, card_title: str, date_str: str) -> str:
    """
    截取指定标题的卡片区域，并返回截图文件路径
    """
    log_info(f"[save_card_screenshot] 开始截取卡片区域，标题={card_title}")

    if not os.path.exists(SNAPSHOT_DIR):
        os.makedirs(SNAPSHOT_DIR)
        log_info(f"[save_card_screenshot] 创建截图目录: {SNAPSHOT_DIR}")

    card_ele = find_visible_card(page, card_title)

    try:
        card_ele.scroll.to_see()
        log_info("[save_card_screenshot] 已滚动到目标卡片")
    except Exception as e:
        log_info(f"[save_card_screenshot] 滚动到目标卡片失败: {e}")

    # 先等原始卡片稳定
    wait_card_ready(page, card_ele, timeout=10)

    safe_title = sanitize_filename(card_title)
    image_name = f"{date_str}_{safe_title}.png"
    image_path = os.path.join(SNAPSHOT_DIR, image_name)

    try:
        clone_card_to_fixed_layer(page, card_ele)

        # 再等克隆层稳定并获取 bbox
        fixed_bbox = get_fixed_clone_bbox(page, timeout=10)

        if fixed_bbox["width"] <= 0 or fixed_bbox["height"] <= 0:
            raise RuntimeError(
                f"[save_card_screenshot] 固定层卡片尺寸异常: "
                f"width={fixed_bbox['width']}, height={fixed_bbox['height']}"
            )

        capture_card_by_cdp(page, fixed_bbox, image_path)

        # 截图后裁掉外圈白边
        trim_white_border(image_path)

        log_info(f"[save_card_screenshot] 卡片截图保存成功，路径={image_path}")
        return image_path

    finally:
        remove_fixed_clone_layer(page)


def save_to_excel(date_str, weekday_str, huawei_employee, external_partner, image_path):
    """
    保存数据到 Excel，并插入图片
    """
    log_info("开始写入 Excel")

    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "统计"

        # 表头
        ws.append(["日期", "星期", "华为员工", "外包员工", "截图"])

    # 当前行号
    row_idx = ws.max_row + 1

    # 写数据
    ws.cell(row=row_idx, column=1, value=date_str)
    ws.cell(row=row_idx, column=2, value=weekday_str)
    ws.cell(row=row_idx, column=3, value=huawei_employee)
    ws.cell(row=row_idx, column=4, value=external_partner)

    # 插入图片
    img = XLImage(image_path)

    # 控制图片大小（可调）
    img.width = 200
    img.height = 120

    # 插入到第5列
    img_cell = f"E{row_idx}"
    ws.add_image(img, img_cell)

    # 调整行高（否则图片显示不全）
    ws.row_dimensions[row_idx].height = 100

    wb.save(EXCEL_FILE)

    log_info("Excel 写入完成")


def open_hr_after_w3_login(browser_page, timeout: int = LOGIN_TIMEOUT):
    """先在 W3 建立公司登录会话，再在同一浏览器的新标签页打开 HR。"""
    log_info("[W3] 打开公司门户，等待自动登录完成")
    browser_page.get(W3_URL)
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            # 检测可见页面文本，避免仅凭加载完成或固定延时判断登录成功。
            logged_in = browser_page.run_js(
                "return location.hostname === 'w3.huawei.com' && "
                "!!document.body && document.body.innerText.includes(arguments[0]);",
                W3_LOGIN_NAME,
                timeout=5,
            )
            if logged_in:
                log_info("[W3] 已检测到用户姓名，登录成功；新开 HR 标签页")
                break
        except Exception:
            # 登录跳转期间页面上下文会短暂失效，下次轮询重新检测。
            pass
        time.sleep(1)
    else:
        raise TimeoutError(
            f"W3 登录等待超过 {timeout} 秒，未在门户页面检测到 {W3_LOGIN_NAME}。"
            "此时尚未打开 HR 页面，请检查本次运行日志。"
        )
    hr_page = browser_page.new_tab()
    hr_page.get(TARGET_URL)
    return hr_page


def run_once() -> None:
    """
    执行一次抓取任务（包含浏览器生命周期）
    """
    log_info("[run_once] 开始执行单次任务")

    browser_page = None
    page = None
    stage = "创建浏览器"
    started = time.monotonic()
    try:
        browser_page = create_page()
        # try:
        #     page.set.window.size(1600, 1000)
        #     log_info("[run_once] 浏览器窗口已设置为 1600x1000")
        # except Exception as e:
        #     log_info(f"[run_once] 设置浏览器窗口大小失败: {e}")

        stage = "W3 自动登录并打开 HR"
        page = open_hr_after_w3_login(browser_page)

        # 等登录
        stage = "等待 HR 员工数量"
        wait_for_login(page, timeout=LOGIN_TIMEOUT)

        log_info("[run_once] 登录检测通过，等待页面卡片稳定渲染")
        time.sleep(3)

        stage = "读取员工数量"
        huawei_employee = get_count_by_label(page, "华为员工")
        external_partner = get_count_by_label(page, "非雇员")

        now = datetime.now()

        csv_date_str = f"{now.year}/{now.month}/{now.day}"
        image_date_str = f"{now.year}-{now.month}-{now.day}"

        weekday_map = {
            0: "周一",
            1: "周二",
            2: "周三",
            3: "周四",
            4: "周五",
            5: "周六",
            6: "周日",
        }
        weekday_str = weekday_map[now.weekday()]

        stage = "截取组织卡片"
        screenshot_path = save_card_screenshot(page, CARD_TITLE, image_date_str)

        stage = "保存 Excel"
        save_to_excel(
            csv_date_str,
            weekday_str,
            huawei_employee,
            external_partner,
            screenshot_path
        )

        log_info(
            f"[run_once] 抓取完成: 日期={csv_date_str}, 星期={weekday_str}, "
            f"华为员工={huawei_employee}, 外包员工={external_partner}, 截图路径={screenshot_path}"
        )

    except BaseException:
        log_info(f"[run_once] 任务失败，阶段={stage}\n{traceback.format_exc()}")
        active_page = page if page is not None else browser_page
        if active_page is not None:
            try:
                address = urlsplit(active_page.url)
                log_info(
                    f"[诊断] 页面标题={active_page.title!r}，"
                    f"页面地址（不含查询参数）={address.scheme}://{address.netloc}{address.path}"
                )
            except Exception as diagnostic_error:
                log_info(f"[诊断] 无法读取页面状态: {diagnostic_error}")
        raise
    finally:
        # 清理异常不能覆盖原始任务异常；成功任务若关闭失败则仍以失败退出。
        task_failed = sys.exc_info()[0] is not None
        if browser_page is not None:
            try:
                browser_page.quit(timeout=10, force=True)
            except Exception:
                log_info(f"[run_once] 浏览器关闭失败\n{traceback.format_exc()}")
                if not task_failed:
                    raise
            else:
                log_info("[run_once] 专用浏览器已关闭，登录配置保留")
        log_info(f"[run_once] 任务收尾完成，耗时={time.monotonic() - started:.1f} 秒")


def wait_for_login(page, timeout: int = 30) -> None:
    """
    等待页面登录完成（通过检测“华为员工”数字是否出现）
    """
    log_info(f"开始等待 HR 员工数量加载（最多 {timeout} 秒）")

    xpath = "xpath://div[contains(@class, 'content')][.//p[normalize-space()='华为员工']]/p[2]"
    start_time = time.time()

    while True:
        try:
            ele = page.ele(xpath, timeout=1)
            text = ele.text.strip() if ele else ''

            if text and text.isdigit():
                log_info(f"检测到登录成功，华为员工={text}")
                return

        except ElementNotFoundError:
            # 元素还没出来，继续等
            pass
        except Exception as e:
            # 其它异常也先记录一下，避免静默失败
            log_info(f"等待登录过程中出现异常: {e}")

        if time.time() - start_time > timeout:
            raise TimeoutError("HR 员工数量加载超时，可能是登录失败、页面出现 500 或元素定位失效；详情请查看本次日志。")

        time.sleep(1)


def get_next_run_time(hour: int = 10, minute: int = 0) -> datetime:
    """
    获取下一次执行时间（每天固定时间）
    """
    now = datetime.now()

    target = now.replace(hour=hour, minute=minute, second=0, microsecond=0)

    if now >= target:
        target += timedelta(days=1)

    return target

def wait_until(target_time: datetime) -> None:
    """
    等待直到目标时间
    """
    now = datetime.now()
    seconds = (target_time - now).total_seconds()

    if seconds > 0:
        log_info(f"距离下一次执行还有 {int(seconds)} 秒")
        time.sleep(seconds)

def main() -> None:
    """
    主函数：执行爬虫
    """
    log_info(f"程序启动，PID={os.getpid()}，Python={sys.executable}，工作目录={os.getcwd()}")
    log_info(f"输出位置：Excel={os.path.abspath(EXCEL_FILE)}，截图={os.path.abspath(SNAPSHOT_DIR)}")
    try:
        run_once()
    except BaseException:
        log_info("程序失败退出（非零退出码）")
        raise
    else:
        log_info("程序正常结束")


if __name__ == "__main__":
    main()
